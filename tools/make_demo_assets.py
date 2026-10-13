#!/usr/bin/env python3
"""
Gera os arquivos de demonstração a partir de database/seeds/content.php:
  - um vídeo curto (MP4/H.264) por aula em vídeo, com título e tópicos da aula;
  - as apostilas em PDF e a planilha de exercícios citadas como materiais.

Requer: php, ffmpeg, reportlab e openpyxl. Rode a partir da raiz do projeto:
    python3 tools/make_demo_assets.py
Os arquivos vão para storage/app/demo/ (já versionados no repositório).
"""
import json
import os
import subprocess
import tempfile
import textwrap

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_VIDEOS = os.path.join(ROOT, "storage/app/demo/videos")
OUT_MATERIALS = os.path.join(ROOT, "storage/app/demo/materials")
FONT_BOLD = os.path.join(ROOT, "resources/fonts/BricolageGrotesque-Bold.ttf")
FONT_BODY = os.path.join(ROOT, "resources/fonts/AtkinsonHyperlegible-Regular.ttf")
W, H = 854, 480


def load_content():
    php = "echo json_encode(require 'database/seeds/content.php', JSON_UNESCAPED_UNICODE);"
    out = subprocess.check_output(["php", "-r", php], cwd=ROOT)
    return json.loads(out)


def write_tmp(text, tmpdir, name):
    path = os.path.join(tmpdir, name)
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)
    return path


def make_video(course, module_idx, lesson_idx, lesson, tmpdir):
    points = lesson.get("points") or []
    duration = 5 + 7 * len(points) + 4
    color = course["color"].lstrip("#")
    name = f"{course['slug']}-m{module_idx}-a{lesson_idx}.mp4"
    out = os.path.join(OUT_VIDEOS, name)

    def fade(start):
        return f"enable='gte(t,{start})':alpha='min(1,(t-{start})/0.6)'"

    filters = ["drawgrid=width=32:height=32:thickness=1:color=white@0.07"]
    course_file = write_tmp(course["title"], tmpdir, f"c{module_idx}{lesson_idx}.txt")
    filters.append(
        f"drawtext=fontfile='{FONT_BODY}':textfile='{course_file}':fontsize=18:fontcolor=white@0.75:x=56:y=44"
    )
    title = "\n".join(textwrap.wrap(lesson["title"], 30))
    title_file = write_tmp(title, tmpdir, f"t{module_idx}{lesson_idx}.txt")
    filters.append(
        f"drawtext=fontfile='{FONT_BOLD}':textfile='{title_file}':fontsize=38:line_spacing=6:fontcolor=white:x=56:y=92:{fade(0.3)}"
    )
    title_lines = title.count("\n") + 1
    y = 92 + title_lines * 46 + 34
    for i, point in enumerate(points):
        start = 3 + i * 7
        filters.append(
            f"drawbox=x=56:y={y + 6}:w=14:h=14:color=0xF5E04B:t=fill:enable='gte(t,{start})'"
        )
        pf = write_tmp(point, tmpdir, f"p{module_idx}{lesson_idx}{i}.txt")
        filters.append(
            f"drawtext=fontfile='{FONT_BODY}':textfile='{pf}':fontsize=24:fontcolor=white:x=86:y={y}:{fade(start)}"
        )
        y += 48
    end_file = write_tmp("Aula de demonstração · conteúdo resumido", tmpdir, f"e{module_idx}{lesson_idx}.txt")
    filters.append(
        f"drawtext=fontfile='{FONT_BODY}':textfile='{end_file}':fontsize=16:fontcolor=white@0.6:x=56:y=h-52"
    )
    graph = (
        f"[0]{','.join(filters)}[bg];"
        f"[bg][1]overlay=x='-w+w*t/{duration}':y=H-6:shortest=1[v]"
    )
    cmd = [
        "ffmpeg", "-y", "-hide_banner", "-loglevel", "error",
        "-f", "lavfi", "-i", f"color=c=0x{color}:s={W}x{H}:d={duration}:r=24",
        "-f", "lavfi", "-i", f"color=c=0xF5E04B:s={W}x6:d={duration}:r=24",
        "-filter_complex", graph, "-map", "[v]",
        "-c:v", "libx264", "-preset", "slow", "-crf", "30", "-tune", "stillimage",
        "-pix_fmt", "yuv420p", "-movflags", "+faststart", out,
    ]
    subprocess.check_call(cmd)
    # Versão WebM (VP9) como fonte alternativa para navegadores sem H.264.
    subprocess.check_call([
        "ffmpeg", "-y", "-hide_banner", "-loglevel", "error", "-i", out,
        "-c:v", "libvpx-vp9", "-crf", "42", "-b:v", "0", "-deadline", "good", "-cpu-used", "4", "-row-mt", "1", "-an",
        out[:-4] + ".webm",
    ])
    return name, duration


def make_pdf(path, title, course_title, sections):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors

    pdfmetrics.registerFont(TTFont("Body", FONT_BODY))
    pdfmetrics.registerFont(TTFont("Bold", FONT_BOLD))
    ink = colors.HexColor("#1F3A93")
    styles = {
        "h1": ParagraphStyle("h1", fontName="Bold", fontSize=22, leading=27, textColor=ink, spaceAfter=4),
        "sub": ParagraphStyle("sub", fontName="Body", fontSize=11, leading=15, textColor=colors.HexColor("#5B6478"), spaceAfter=18),
        "h2": ParagraphStyle("h2", fontName="Bold", fontSize=14, leading=18, spaceBefore=12, spaceAfter=6),
        "p": ParagraphStyle("p", fontName="Body", fontSize=11, leading=17),
        "li": ParagraphStyle("li", fontName="Body", fontSize=11, leading=17, leftIndent=12, bulletIndent=0),
    }
    story = [Paragraph(title, styles["h1"]), Paragraph(course_title + " · material de apoio", styles["sub"])]
    for heading, items in sections:
        story.append(Paragraph(heading, styles["h2"]))
        for item in items:
            if item.startswith("- "):
                story.append(Paragraph(item[2:], styles["li"], bulletText="•"))
            else:
                story.append(Paragraph(item, styles["p"]))
                story.append(Spacer(1, 4))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Body", 8.5)
        canvas.setFillColor(colors.HexColor("#5B6478"))
        canvas.drawString(20 * mm, 12 * mm, "Material de demonstração · Instituto Horizonte (instituição fictícia)")
        canvas.drawRightString(190 * mm, 12 * mm, f"Página {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=22 * mm, bottomMargin=22 * mm,
                            title=title, author="Instituto Horizonte (demo)")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def make_materials():
    os.makedirs(OUT_MATERIALS, exist_ok=True)
    make_pdf(os.path.join(OUT_MATERIALS, "excel-apostila-modulo-1.pdf"), "Primeiros passos no Excel", "Excel Essencial para o Trabalho", [
        ("1. Organização dos dados", [
            "Uma planilha bem estruturada é a base de qualquer análise. Antes de pensar em fórmulas, garanta que os dados estejam em formato de lista.",
            "- Títulos de coluna na primeira linha, sem células mescladas",
            "- Um registro por linha e um tipo de informação por coluna",
            "- Nada de linhas ou colunas em branco no meio dos dados",
        ]),
        ("2. Tipos de dado", [
            "O Excel trata texto, números e datas de forma diferente. Um número digitado como texto não entra em somas; uma data digitada como texto não pode ser filtrada por mês.",
            "- Números alinham à direita por padrão; textos, à esquerda",
            "- Use a formatação de moeda em vez de digitar o símbolo R$",
        ]),
        ("Exercício 1", [
            "Monte uma lista com 10 despesas da sua casa contendo: data, descrição, categoria e valor. Formate os valores como moeda e as datas no formato dd/mm/aaaa.",
        ]),
    ])
    make_pdf(os.path.join(OUT_MATERIALS, "atendimento-guia-mensagens.pdf"), "Guia de mensagens prontas", "Atendimento ao Cliente e Comunicação", [
        ("Primeiro contato", [
            "Olá, Maria! Aqui é o Paulo, da equipe de atendimento. Recebi sua mensagem e já estou verificando o seu pedido.",
        ]),
        ("Quando precisa de mais tempo", [
            "Maria, ainda estou aguardando o retorno do setor responsável. Te dou uma posição até amanhã às 10h.",
        ]),
        ("Recusa com alternativa", [
            "Não conseguimos trocar o produto depois de 30 dias, porque o prazo de troca já passou. Posso abrir um chamado de garantia com o fabricante, que costuma responder em até 5 dias úteis. Quer que eu faça?",
        ]),
        ("Encerramento", [
            "Pronto, Maria: o reembolso foi registrado e cai na fatura em até 2 ciclos. Se precisar, é só chamar por aqui.",
        ]),
    ])
    make_pdf(os.path.join(OUT_MATERIALS, "lgpd-checklist-adequacao.pdf"), "Checklist de adequação à LGPD", "LGPD na Prática para Empresas", [
        ("Governança", [
            "- Encarregado (DPO) indicado e com contato publicado",
            "- Política de privacidade clara e atualizada",
            "- Treinamento da equipe que lida com dados pessoais",
        ]),
        ("Mapeamento", [
            "- Inventário de dados por setor: quais dados, de quem, para quê, onde ficam",
            "- Base legal definida para cada finalidade",
            "- Prazos de retenção e descarte definidos",
        ]),
        ("Direitos do titular", [
            "- Canal para pedidos de acesso, correção e eliminação",
            "- Procedimento interno com prazos de resposta",
        ]),
        ("Segurança", [
            "- Controle de acesso por perfil e senhas fortes",
            "- Backups testados periodicamente",
            "- Plano de resposta a incidentes e registro de ocorrências",
        ]),
    ])

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"
    head = ["Data", "Vendedor", "Produto", "Quantidade", "Preço unitário", "Total"]
    ws.append(head)
    rows = [
        ("2026-03-02", "Ana", "Caderno", 12, 18.9), ("2026-03-03", "Bruno", "Caneta", 40, 3.5),
        ("2026-03-05", "Ana", "Mochila", 3, 149.0), ("2026-03-08", "Carla", "Caderno", 7, 18.9),
        ("2026-03-11", "Bruno", "Estojo", 9, 24.0), ("2026-03-14", "Carla", "Mochila", 2, 149.0),
        ("2026-03-18", "Ana", "Caneta", 25, 3.5), ("2026-03-21", "Bruno", "Caderno", 15, 18.9),
    ]
    import datetime
    for i, (d, v, p, q, pu) in enumerate(rows, start=2):
        ws.append([datetime.date.fromisoformat(d), v, p, q, pu, f"=D{i}*E{i}"])
        ws[f"A{i}"].number_format = "DD/MM/YYYY"
        ws[f"E{i}"].number_format = '"R$" #,##0.00'
        ws[f"F{i}"].number_format = '"R$" #,##0.00'
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F7A4D")
    for col, width in zip("ABCDEF", (12, 12, 12, 12, 16, 14)):
        ws.column_dimensions[col].width = width
    ex = wb.create_sheet("Exercícios")
    for line in [
        "1. Na coluna G da aba Vendas, use =SE(F2>200;\"Meta\";\"Abaixo\") e copie para baixo.",
        "2. Em uma célula livre, calcule o total vendido com =SOMA(F2:F9).",
        "3. Calcule a média de quantidade com =MÉDIA(D2:D9).",
        "4. Crie uma aba Preços com produto e preço e use PROCV para buscar o preço do Caderno.",
    ]:
        ex.append([line])
    ex.column_dimensions["A"].width = 100
    wb.save(os.path.join(OUT_MATERIALS, "excel-exercicios-formulas.xlsx"))


def main():
    os.makedirs(OUT_VIDEOS, exist_ok=True)
    content = load_content()
    durations = {}
    with tempfile.TemporaryDirectory() as tmp:
        for course in content["courses"]:
            for mi, module in enumerate(course["modules"], start=1):
                for li, lesson in enumerate(module["lessons"], start=1):
                    if lesson.get("type") != "video":
                        continue
                    name, duration = make_video(course, mi, li, lesson, tmp)
                    durations[name] = duration
                    print("vídeo", name, f"{duration}s")
    with open(os.path.join(OUT_VIDEOS, "durations.json"), "w") as f:
        json.dump(durations, f, indent=1)
    make_materials()
    print("materiais ok")


if __name__ == "__main__":
    main()
